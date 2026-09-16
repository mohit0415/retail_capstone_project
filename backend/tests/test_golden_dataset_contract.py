"""HTTP contract check for every row of the 50-record golden test dataset.

Drives each question from ``evals/test_dataset_50.xlsx`` through ``POST /ask``
with a token minted for that row's role, and asserts the status-code contract
the dataset promises:

  * ``pending_review`` rows come back **202** with a poll_url and no answer
  * every other row (``answered`` / ``refused`` / ``clarification_required`` /
    ``refused_or_scoped``) comes back **200** with the matching body status

The graph is stubbed per row (same FakeGraph pattern as test_api_routes.py),
so this is offline and deterministic: it proves the API maps each expected
terminal outcome for each dataset role to the right HTTP response. Answer
quality against the live model is the job of ``evals/run_api_eval.py``.
"""

from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api import routes
from src.cache import response_cache as response_cache_module
from src.cache.response_cache import ResponseCache
from src.core import conversation
from src.core.idempotency import IdempotencyStore
from src.schemas.enums import EvidencePath, RiskLevel, Role, TerminalOutcome
from src.schemas.models import ConfidenceBreakdown, DraftAnswer, RetrievedChunk, RiskAssessment

DATASET = Path(__file__).resolve().parents[1] / "evals" / "test_dataset_50.xlsx"

EXPECTED_HTTP = {
    "answered": 200,
    "refused": 200,
    "clarification_required": 200,
    "refused_or_scoped": 200,
    "pending_review": 202,
}


def load_rows() -> list[dict]:
    ws = load_workbook(DATASET, read_only=True)["Test Dataset"]
    rows = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue

        rows.append(
            {
                "test_id": str(row[1]),
                "question": str(row[3]),
                "role": Role(str(row[4])),
                "expected_outcome": str(row[8]),
            }
        )

    return rows


ROWS = load_rows()


def test_the_dataset_has_the_mandated_50_rows():
    assert len(ROWS) == 50


@pytest.fixture
def client(no_db, no_tracing, auth0, monkeypatch):
    from main import app

    monkeypatch.setattr(routes, "idempotency_store", IdempotencyStore())
    monkeypatch.setattr(response_cache_module, "response_cache", ResponseCache(embedder=None, semantic_threshold=1.0))
    monkeypatch.setattr(conversation, "record_exchange", lambda *args, **kwargs: None)
    monkeypatch.setattr(conversation, "record_user_turn", lambda *args, **kwargs: None)
    monkeypatch.setattr(conversation, "record_assistant_turn", lambda *args, **kwargs: None)
    monkeypatch.setattr(conversation, "load_thread", lambda thread_id: ([], ""))
    monkeypatch.setattr(routes.slo, "record_latency", lambda *args, **kwargs: None)

    return TestClient(app)


class FakeGraph:
    def __init__(self, final_state: dict):
        self.final_state = final_state

    def invoke(self, state, config=None):
        return {**state, **self.final_state}


def _chunk() -> RetrievedChunk:
    return RetrievedChunk(
        chunk_id="retention-4.2",
        doc_type="retention_policy",
        document_title="Data Retention & Archival Policy",
        section="Retention Period Standards",
        clause_number="4.2",
        version="2.0",
        content="Customer transaction records are retained for seven years.",
    )


def state_for_outcome(expected_outcome: str) -> dict:
    if expected_outcome == "answered":
        return {
            "terminal_outcome": TerminalOutcome.ANSWERED.value,
            "draft": DraftAnswer(
                answer="Seven years [Data Retention & Archival Policy §4.2].",
                cited_clauses=["Data Retention & Archival Policy §4.2"],
            ),
            "retrieved_chunks": [_chunk()],
            "confidence": ConfidenceBreakdown(final_score=0.9),
            "risk": RiskAssessment(final_level=RiskLevel.LOW),
            "evidence_path": EvidencePath.RAG.value,
        }

    if expected_outcome == "pending_review":
        return {
            "terminal_outcome": TerminalOutcome.ESCALATED.value,
            "escalation_reason": "risk level is High",
            "risk": RiskAssessment(final_level=RiskLevel.HIGH),
        }

    if expected_outcome == "clarification_required":
        return {
            "terminal_outcome": TerminalOutcome.CLARIFICATION_REQUIRED.value,
            "clarification_question": "Which vendor do you mean?",
        }

    # refused and refused_or_scoped
    return {
        "terminal_outcome": TerminalOutcome.REFUSED.value,
        "refusal_reason": "request rejected by the input guardrail",
    }


@pytest.mark.parametrize("row", ROWS, ids=[r["test_id"] for r in ROWS])
def test_each_golden_row_gets_the_promised_http_status(row, client, auth0, monkeypatch):
    monkeypatch.setattr(routes, "get_compiled_graph", lambda: FakeGraph(state_for_outcome(row["expected_outcome"])))

    token = auth0(row["role"], user_id=f"eval-{row['role'].value}")
    response = client.post(
        "/ask",
        json={"query": row["question"]},
        headers={"Authorization": f"Bearer {token}"},
    )

    expected_status = EXPECTED_HTTP[row["expected_outcome"]]
    assert response.status_code == expected_status, (
        f"{row['test_id']}: expected {expected_status}, got {response.status_code}: {response.text[:200]}"
    )

    body = response.json()

    if row["expected_outcome"] == "pending_review":
        assert body["status"] == "pending_review"
        assert body["poll_url"]
        assert "answer" not in body
    elif row["expected_outcome"] == "refused_or_scoped":
        assert body["status"] in ("refused", "answered")
    else:
        assert body["status"] == row["expected_outcome"]

    if body["status"] == "answered":
        assert body["answer"]
        assert body["citations"]
