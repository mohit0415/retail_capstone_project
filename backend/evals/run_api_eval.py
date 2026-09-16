"""Run the 50-record golden test dataset against the live /ask API.

Reads the questions from evals/test_dataset_50.xlsx (and the reference answers
from evals/test_dataset_50_with_answers.xlsx), POSTs each one to {base-url}/ask
with a real bearer token, and checks the HTTP contract:

  * rows whose Expected Outcome is ``pending_review``       -> HTTP 202
  * every other row (answered / refused / clarification)    -> HTTP 200
  * ``refused_or_scoped``                                   -> HTTP 200 with
    body status ``refused`` or ``answered``

The body ``status`` field is compared against the row's Expected Outcome, the
returned answer is scored against the reference answer (keyword coverage +
sequence similarity - a cheap proxy, not an LLM judge), and everything lands in
an Excel report next to the dataset plus a JSON copy for diffing between runs.

Auth: the API only trusts Auth0 tokens, so grab one from a logged-in frontend
session (browser dev tools -> the Authorization header on any /ask call) and
export it:

    export EVAL_TOKEN="eyJ..."

The role encoded in the token is what the API enforces. Rows for other roles
are still sent (the graph's RBAC reaction is itself worth observing), but the
report marks rows whose dataset role differs from --token-role so a mismatch
is never mistaken for a model failure. Per-role tokens can be supplied as
EVAL_TOKEN_STORE_MANAGER, EVAL_TOKEN_STORE_ASSOCIATE, ... and are preferred
over EVAL_TOKEN when present.

Usage:
    uv run python evals/run_api_eval.py                       # localhost:8000
    uv run python evals/run_api_eval.py --base-url https://rpids-api.onrender.com
    uv run python evals/run_api_eval.py --filter hr-          # one category
"""

import argparse
import difflib
import json
import os
import re
import statistics
import sys
import time
import uuid
from dataclasses import asdict, dataclass, field
from pathlib import Path

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

EVALS_DIR = Path(__file__).parent
DATASET = EVALS_DIR / "test_dataset_50.xlsx"
ANSWERS = EVALS_DIR / "test_dataset_50_with_answers.xlsx"

# Expected Outcome -> HTTP status the API contract promises for it.
STATUS_FOR_OUTCOME = {
    "answered": 200,
    "refused": 200,
    "clarification_required": 200,
    "refused_or_scoped": 200,
    "pending_review": 202,
}

STOPWORDS = frozenset(
    "the a an and or of to in for is are was were be been with on at by from as it its this that "
    "what which who whom how when where do does did can could should would may must not no".split()
)


@dataclass
class Case:
    test_id: str
    category: str
    question: str
    role: str
    expected_outcome: str
    reference_answer: str = ""


@dataclass
class CaseResult:
    test_id: str
    category: str
    question: str
    role: str
    token_role_mismatch: bool
    expected_outcome: str
    expected_status: int
    http_status: int | None = None
    body_status: str = ""
    status_ok: bool = False
    outcome_ok: bool = False
    answer: str = ""
    reference_answer: str = ""
    keyword_coverage: float | None = None
    similarity: float | None = None
    confidence: float | None = None
    risk_level: str = ""
    evidence_path: str = ""
    citations: int = 0
    poll_url: str = ""
    elapsed_seconds: float = 0.0
    error: str = ""


@dataclass
class RunReport:
    base_url: str
    token_role: str
    results: list[CaseResult] = field(default_factory=list)


def load_cases(filter_prefix: str) -> list[Case]:
    ws = load_workbook(DATASET, read_only=True)["Test Dataset"]
    cases: dict[str, Case] = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue

        _, test_id, category, question, role, _intent, _path, _risk, outcome, *_ = row
        cases[str(test_id)] = Case(
            test_id=str(test_id),
            category=str(category),
            question=str(question),
            role=str(role),
            expected_outcome=str(outcome),
        )

    answers_ws = load_workbook(ANSWERS, read_only=True)["QA Test Dataset"]

    for row in answers_ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue

        test_id, reference = str(row[1]), row[6]

        if test_id in cases and reference:
            cases[test_id].reference_answer = str(reference)

    ordered = list(cases.values())

    if filter_prefix:
        ordered = [c for c in ordered if c.test_id.startswith(filter_prefix)]

    return ordered


def token_for_role(role: str, default_token: str) -> str:
    return os.environ.get(f"EVAL_TOKEN_{role.upper()}", default_token)


def keywords(text: str) -> set[str]:
    return {w for w in re.findall(r"[a-z0-9]+", text.lower()) if len(w) > 2 and w not in STOPWORDS}


def score_answer(answer: str, reference: str) -> tuple[float, float]:
    """(keyword coverage of the reference in the answer, sequence similarity)."""
    ref_kw = keywords(reference)
    coverage = len(ref_kw & keywords(answer)) / len(ref_kw) if ref_kw else 0.0
    similarity = difflib.SequenceMatcher(None, answer.lower(), reference.lower()).ratio()

    return round(coverage, 3), round(similarity, 3)


def run_case(client: httpx.Client, case: Case, token: str, token_role: str) -> CaseResult:
    result = CaseResult(
        test_id=case.test_id,
        category=case.category,
        question=case.question,
        role=case.role,
        token_role_mismatch=case.role != token_role
        and f"EVAL_TOKEN_{case.role.upper()}" not in os.environ,
        expected_outcome=case.expected_outcome,
        expected_status=STATUS_FOR_OUTCOME.get(case.expected_outcome, 200),
        reference_answer=case.reference_answer,
    )

    started = time.monotonic()

    try:
        response = client.post(
            "/ask",
            json={"query": case.question},
            headers={
                "Authorization": f"Bearer {token}",
                # unique per run so the idempotency store never replays an old verdict
                "Idempotency-Key": f"eval-{case.test_id}-{uuid.uuid4().hex[:12]}",
            },
        )
    except httpx.HTTPError as exc:
        result.elapsed_seconds = round(time.monotonic() - started, 3)
        result.error = f"{type(exc).__name__}: {exc}"[:300]

        return result

    result.elapsed_seconds = round(time.monotonic() - started, 3)
    result.http_status = response.status_code

    try:
        body = response.json()
    except ValueError:
        body = {}

    result.body_status = str(body.get("status", ""))
    result.status_ok = response.status_code == result.expected_status

    if case.expected_outcome == "refused_or_scoped":
        result.outcome_ok = result.body_status in ("refused", "answered")
    else:
        result.outcome_ok = result.body_status == case.expected_outcome

    result.answer = str(body.get("answer") or body.get("reason") or body.get("detail") or "")
    result.confidence = body.get("confidence")
    result.risk_level = str(body.get("risk_level", ""))
    result.evidence_path = str(body.get("evidence_path", ""))
    result.citations = len(body.get("citations") or [])
    result.poll_url = str(body.get("poll_url", ""))

    if result.answer and case.reference_answer:
        result.keyword_coverage, result.similarity = score_answer(result.answer, case.reference_answer)

    return result


def percentile(values: list[float], p: float) -> float:
    if not values:
        return 0.0

    ordered = sorted(values)

    return ordered[min(len(ordered) - 1, round(p * (len(ordered) - 1)))]


def summarise(report: RunReport) -> dict:
    results = report.results
    reached = [r for r in results if r.http_status is not None]
    scored = [r for r in results if r.keyword_coverage is not None]
    latencies = [r.elapsed_seconds for r in reached]

    by_category: dict[str, dict] = {}

    for r in results:
        bucket = by_category.setdefault(r.category, {"total": 0, "status_ok": 0, "outcome_ok": 0})
        bucket["total"] += 1
        bucket["status_ok"] += r.status_ok
        bucket["outcome_ok"] += r.outcome_ok

    return {
        "base_url": report.base_url,
        "cases_run": len(results),
        "transport_errors": len(results) - len(reached),
        "http_status_pass": sum(r.status_ok for r in results),
        "http_status_pass_rate": round(sum(r.status_ok for r in results) / len(results), 4) if results else 0.0,
        "outcome_match": sum(r.outcome_ok for r in results),
        "outcome_match_rate": round(sum(r.outcome_ok for r in results) / len(results), 4) if results else 0.0,
        "answered_with_citations": sum(1 for r in results if r.body_status == "answered" and r.citations > 0),
        "answered_total": sum(1 for r in results if r.body_status == "answered"),
        "mean_keyword_coverage": round(statistics.mean(r.keyword_coverage for r in scored), 3) if scored else None,
        "mean_similarity": round(statistics.mean(r.similarity for r in scored), 3) if scored else None,
        "latency_p50": round(statistics.median(latencies), 3) if latencies else 0.0,
        "latency_p95": round(percentile(latencies, 0.95), 3) if latencies else 0.0,
        "by_category": by_category,
    }


PASS_FILL = PatternFill("solid", start_color="C6EFCE")
FAIL_FILL = PatternFill("solid", start_color="FFC7CE")

RESULT_COLUMNS = [
    ("Test ID", 10),
    ("Category", 16),
    ("Role", 18),
    ("Question", 60),
    ("Expected Outcome", 18),
    ("Expected HTTP", 14),
    ("Actual HTTP", 12),
    ("Status Pass", 11),
    ("Body Status", 20),
    ("Outcome Match", 14),
    ("Answer", 80),
    ("Reference Answer", 80),
    ("Keyword Coverage", 16),
    ("Similarity", 11),
    ("Confidence", 11),
    ("Risk", 10),
    ("Evidence Path", 14),
    ("Citations", 10),
    ("Latency (s)", 11),
    ("Role/Token Mismatch", 18),
    ("Error", 40),
]


def write_excel(report: RunReport, summary: dict, out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    for col, (header, width) in enumerate(RESULT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = width

    for row_idx, r in enumerate(report.results, start=2):
        passed = r.status_ok and r.outcome_ok
        values = [
            r.test_id, r.category, r.role, r.question,
            r.expected_outcome, r.expected_status, r.http_status,
            "PASS" if r.status_ok else "FAIL",
            r.body_status,
            "PASS" if r.outcome_ok else "FAIL",
            r.answer[:1500], r.reference_answer[:1500],
            r.keyword_coverage, r.similarity, r.confidence,
            r.risk_level, r.evidence_path, r.citations,
            r.elapsed_seconds,
            "yes" if r.token_role_mismatch else "",
            r.error,
        ]

        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=value)

        ws.cell(row=row_idx, column=8).fill = PASS_FILL if r.status_ok else FAIL_FILL
        ws.cell(row=row_idx, column=10).fill = PASS_FILL if r.outcome_ok else FAIL_FILL
        ws.cell(row=row_idx, column=1).fill = PASS_FILL if passed else FAIL_FILL

    ws.freeze_panes = "A2"

    summary_ws = wb.create_sheet("Summary")
    summary_ws.column_dimensions["A"].width = 34
    summary_ws.column_dimensions["B"].width = 40

    row = 1

    for key, value in summary.items():
        if key == "by_category":
            continue

        summary_ws.cell(row=row, column=1, value=key).font = Font(bold=True)
        summary_ws.cell(row=row, column=2, value=value if not isinstance(value, dict) else json.dumps(value))
        row += 1

    row += 1
    summary_ws.cell(row=row, column=1, value="Category").font = Font(bold=True)
    summary_ws.cell(row=row, column=2, value="status ok / outcome ok / total").font = Font(bold=True)

    for category, bucket in summary["by_category"].items():
        row += 1
        summary_ws.cell(row=row, column=1, value=category)
        summary_ws.cell(
            row=row, column=2,
            value=f"{bucket['status_ok']} / {bucket['outcome_ok']} / {bucket['total']}",
        )

    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Run the 50-question golden dataset against /ask")
    parser.add_argument("--base-url", default=os.environ.get("EVAL_BASE_URL", "http://localhost:8000"))
    parser.add_argument("--token", default=os.environ.get("EVAL_TOKEN", ""))
    parser.add_argument(
        "--token-role",
        default=os.environ.get("EVAL_TOKEN_ROLE", "compliance_officer"),
        help="role carried by --token, used to flag rows written for a different role",
    )
    parser.add_argument("--filter", default="", help="only run test ids with this prefix (pl-, rl-, cc-, hr-, adv-)")
    parser.add_argument("--timeout", type=float, default=90.0, help="per-request timeout in seconds")
    parser.add_argument("--out", default=str(EVALS_DIR / "api_eval_results.xlsx"))
    parser.add_argument("--ci", action="store_true", help="exit non-zero when any status/outcome check fails")
    args = parser.parse_args()

    if not args.token:
        sys.exit(
            "No token. Log in on the frontend, copy the Authorization bearer token "
            "from any API call in the browser dev tools, then: export EVAL_TOKEN='eyJ...'"
        )

    cases = load_cases(args.filter)

    if not cases:
        sys.exit(f"no cases match filter {args.filter!r}")

    print(f"{len(cases)} cases -> {args.base_url}/ask\n")

    report = RunReport(base_url=args.base_url, token_role=args.token_role)

    with httpx.Client(base_url=args.base_url, timeout=args.timeout) as client:
        for case in cases:
            result = run_case(client, case, token_for_role(case.role, args.token), args.token_role)
            report.results.append(result)

            flag = "PASS" if result.status_ok and result.outcome_ok else "FAIL"
            print(
                f"{result.test_id:8s} {flag}  http={result.http_status} "
                f"(want {result.expected_status})  status={result.body_status or '-':22s} "
                f"{result.elapsed_seconds:6.2f}s  {result.error}"
            )

    summary = summarise(report)

    out_xlsx = Path(args.out)
    write_excel(report, summary, out_xlsx)

    out_json = out_xlsx.with_suffix(".json")
    out_json.write_text(
        json.dumps({"summary": summary, "results": [asdict(r) for r in report.results]}, indent=2),
        encoding="utf-8",
    )

    print("\n" + json.dumps(summary, indent=2))
    print(f"\nreport: {out_xlsx}\n        {out_json}")

    if args.ci and (summary["http_status_pass"] < summary["cases_run"] or summary["outcome_match"] < summary["cases_run"]):
        sys.exit(1)


if __name__ == "__main__":
    main()
